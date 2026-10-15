"""Importer regression: native lists, blank image paragraphs and source order.

Run with python-docx installed: python tests/medical-hub-docx-import-test.py
"""
import base64
import copy
import hashlib
import importlib.util
import io
import tempfile
from pathlib import Path
from docx import Document
from docx.oxml import parse_xml
from docx.oxml.ns import qn

spec = importlib.util.spec_from_file_location('source_import', Path(__file__).parents[1] / 'scripts/import-medical-hub-docx.py')
module = importlib.util.module_from_spec(spec)
spec.loader.exec_module(module)
png = base64.b64decode('iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAwMCAO+j2ioAAAAASUVORK5CYII=')
doc = Document()
numbering = doc.part.numbering_part.element
namespace = 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'
for num_id, fmt, pattern, start in [(7001, 'bullet', '●', 1), (7002, 'decimal', '%1.', 1), (7003, 'decimal', '%1.', 5)]:
    numbering.append(parse_xml(f'<w:abstractNum xmlns:w="{namespace}" w:abstractNumId="{num_id}"><w:lvl w:ilvl="0"><w:start w:val="{start}"/><w:numFmt w:val="{fmt}"/><w:lvlText w:val="{pattern}"/></w:lvl></w:abstractNum>'))
    numbering.append(parse_xml(f'<w:num xmlns:w="{namespace}" w:numId="{num_id}"><w:abstractNumId w:val="{num_id}"/></w:num>'))

def native(text, num_id):
    p = doc.add_paragraph(text)
    props = p._p.get_or_add_pPr().get_or_add_numPr()
    props.get_or_add_numId().val = num_id
    props.get_or_add_ilvl().val = 0
    return p

doc.add_heading('KAPITULLI 11 • TEST', 1)
doc.add_paragraph('Chapter introduction')
doc.add_paragraph('Shënim redaktorial: Source annotation').runs[0].bold = True
doc.add_heading('1. FIRST', 2)
doc.add_paragraph('Lesson opening body')
doc.add_heading('RX • FIRST', 3)
native('Native bullet', 7001).runs[0].bold = True
native('First numbered item', 7002)
native('Second numbered item', 7002)
doc.add_paragraph().add_run().add_picture(io.BytesIO(png))
doc.add_paragraph('After image')
table = doc.add_table(rows=2, cols=1)
table.cell(0, 0).text = 'Header'
table.cell(1, 0).text = 'Cell'
doc.add_heading('2. RX • SECOND', 2)
doc.add_heading('HAPAT PRAKTIKË', 3)
native('Continued numbered item', 7002)
native('Restarted at five', 7003)
doc.add_heading('BURIME PËR SHËNIMET REDAKTORIALE', 3)
doc.add_paragraph('Source reference')
note = 'Titulli dhe renditja vijnë nga libri kryesor. Përmbajtja klinike duhet të shkruhet dhe verifikohet në Studio para publikimit.'
existing = [{'_id': f'fixture-topic-{order}', '_rev': 'fixture-revision', 'chapter': {'_ref': 'fixture-chapter-11'}, 'order': order, 'sections': [{'_key': 'editorial-start', 'content': [{'_key': 'editorial-note', '_type': 'clinicalCallout', 'body': [{'children': [{'text': note}]}]}]}]} for order in [1, 2]]
assets = {hashlib.sha256(png).hexdigest(): {'assetId': 'image-fixture-1x1-png'}}
with tempfile.TemporaryDirectory() as temp:
    path = Path(temp) / 'fixture.docx'
    doc.save(path)
    try:
        module.build(path, existing, True)
    except AssertionError as error:
        assert 'Upload source image' in str(error)
    else:
        raise AssertionError('A blank image paragraph was silently dropped')
    result = module.build(path, existing, True, assets)
    intro = result['topics'][0]['sections'][0]
    assert intro['title'] == '__SOURCE_BODY__' and intro['_key'] == 'source-11-intro'
    assert [''.join(s['text'] for s in b['children']) for b in intro['content']] == ['Chapter introduction', 'Shënim redaktorial: Source annotation']
    assert intro['content'][1]['children'][0]['marks'] == ['strong']
    assert result['counts']['introductionBlocks'] == 2
    assert len(result['topics'][1]['sections']) == 2, 'Do not repeat the introduction in subsequent lessons'
    rebuilt = module.build(path, [{**existing[i], 'originalTitle': t['originalTitle'], 'sections': t['sections']} for i, t in enumerate(result['topics'])], False, assets)
    assert [t['sections'] for t in rebuilt['topics']] == [t['sections'] for t in result['topics']], 'Introduction import must be repeatable'
    assert result['topics'][0]['sections'][1]['content'][0]['children'][0]['text'] == 'Lesson opening body'
    keys = [s['_key'] for s in result['topics'][0]['sections']]
    assert len(keys) == len(set(keys)), 'Chapter introduction and lesson opening need separate identities'
    first = result['topics'][0]['sections'][2]['content']
    assert [b['_type'] for b in first] == ['block', 'block', 'block', 'medicalFigure', 'block', 'medicalTable']
    assert first[0]['children'][0]['text'] == '● '
    assert first[0]['children'][1]['marks'] == ['strong']
    assert first[3]['image']['asset']['_ref'] == assets[hashlib.sha256(png).hexdigest()]['assetId']
    second_volume_id = '1QNVPfGpPp3lghWIcPf3ynT1rrFMTA5IilJuZ1tEtI_M'
    second_volume = module.build(path, existing, True, assets, second_volume_id)
    assert second_volume['sourceDocumentId'] == second_volume_id
    assert second_volume['sourceSha256'] == result['sourceSha256']
    second_figure = second_volume['topics'][0]['sections'][2]['content'][3]
    assert second_figure['sourceUrl'] == f'https://docs.google.com/document/d/{second_volume_id}/edit'
    assert second_figure['image'] == first[3]['image'], 'Changing volume attribution must preserve the source asset'
    second = result['topics'][1]['sections'][0]['content']
    assert [b['children'][0]['text'] for b in second] == ['3. ', '5. ']
    assert [s['sectionType'] for s in result['topics'][1]['sections']] == ['prescription', 'general']
    assert result['counts']['nativeListItems'] == 5
    edited = copy.deepcopy(existing)
    edited[0]['sections'][0]['content'][0]['body'][0]['children'][0]['text'] = 'Authored note'
    try:
        module.build(path, edited, True, assets)
    except AssertionError as error:
        assert 'Edited migration note' in str(error)
    else:
        raise AssertionError('An authored note was overwritten')
with tempfile.TemporaryDirectory() as temp:
    dotted = Document()
    dotted.add_heading('KAPITULLI 16 • TEST', 1)
    dotted.add_paragraph('Chapter introduction')
    dotted.add_heading('16.1 • Lesson title', 2)
    dotted.add_heading('RX • TREATMENT', 3)
    dotted.add_paragraph('Original treatment').runs[0].bold = True
    dotted.add_heading('Referencat e shënimeve klinike — kapitulli 16', 2)
    dotted.add_paragraph('Original reference')
    path = Path(temp) / 'dotted.docx'
    dotted.save(path)
    scaffold = copy.deepcopy(existing[0])
    scaffold.update(_id='fixture-topic-16-1', chapter={'_ref': 'fixture-chapter-16'})
    result = module.build(path, [scaffold], True)
    topic = result['topics'][0]
    assert topic['originalTitle'] == '16.1 • Lesson title'
    assert topic['title'] == 'Lesson title'
    assert [s['sectionType'] for s in topic['sections']] == ['general', 'prescription', 'general']
    assert topic['sections'][-1]['title'] == 'Referencat e shënimeve klinike — kapitulli 16'
    assert topic['sections'][-1]['content'][0]['children'][0]['text'] == 'Original reference'
    assert topic['sections'][1]['content'][0]['children'][0]['marks'] == ['strong']
    try:
        module.lesson_heading('15.1 • Wrong chapter', 16)
    except AssertionError:
        pass
    else:
        raise AssertionError('A mismatched dotted chapter was accepted')
with tempfile.TemporaryDirectory() as temp:
    continued = Document()
    continued.add_heading('KAPITULLI 24 — TEST', 1)
    continued.add_heading('24.1 — Antibiotikët', 2)
    continued.add_paragraph('Original opening paragraph')
    for text in ['First table', 'Second table']:
        continued.add_heading('Antibiotikët — vijim', 3)
        table = continued.add_table(rows=2, cols=1)
        table.cell(0, 0).text = 'Header'
        table.cell(1, 0).text = text
    path = Path(temp) / 'continued.docx'
    continued.save(path)
    scaffold = copy.deepcopy(existing[0])
    scaffold.update(_id='fixture-topic-24-1', chapter={'_ref': 'fixture-chapter-24'})
    result = module.build(path, [scaffold], True)
    topic = result['topics'][0]
    assert topic['originalTitle'] == '24.1 — Antibiotikët'
    assert topic['title'] == 'Antibiotikët'
    assert [s['title'] for s in topic['sections']] == ['__SOURCE_BODY__', 'Antibiotikët — vijim', 'Antibiotikët — vijim']
    assert len({s['_key'] for s in topic['sections']}) == 3, 'Repeated continuation headings and opening body need distinct identities'
    assert [s['content'][0]['rows'][0]['cells'][0] for s in topic['sections'][1:]] == ['First table', 'Second table']
    rebuilt = module.build(path, [{**scaffold, 'originalTitle': topic['originalTitle'], 'sections': topic['sections']}])
    assert rebuilt['topics'][0]['sections'] == topic['sections'], 'Repeated source sections must remain distinct on rebuild'
print('DOCX import: source lists, bold, images/order, author guard, lesson delimiters, repeated continuation sections and trailing references passed.')

# The prescription mode uses its own canonical document and fails closed when
# the source paragraphs change. No clinical content is inferred from fragments.
repair_spec = importlib.util.spec_from_file_location('rx_repairs', Path(__file__).parents[1] / 'scripts/repair-prescription-source.py')
repairs = importlib.util.module_from_spec(repair_spec)
repair_spec.loader.exec_module(repairs)
with tempfile.TemporaryDirectory() as temp:
    rx_source = Document()
    rx_source.add_heading('KAPITULLI 5 - QA',2)
    sample = []
    for index,ident in enumerate([repairs.DENGUE,repairs.COVID,repairs.PLEURAL]+[f'qa-{i}' for i in range(71)]):
        heading = f'SOURCE QA {index}'
        rx_source.add_heading(heading,4)
        sample.append({'_id':ident,'_rev':'qa-rev','reviewStatus':'verified','sourceHeading':heading,'logicBlocks':[]})
    by_id = {g['_id']:g for g in sample}
    def b(key,sig='',title='QA'):
        return {'_key':key,'_type':'prescriptionLogicBlock','order':int(key[1:]),'items':[{'_key':'i'+key[1:],'title':title,'sig':sig}]}
    vitamin = '2 mL/ampulë, IV ngadalë, OD.'
    remdesivir = 'Infuzion IV në 100 mL 0.9% NaCl si dozë ngarkuese; pastaj remdesivir 100 mg IV në 100 mL 0.9% NaCl, OD, për 45 minuta deri 1 orë. Gjithsej 4-9 doza nga 100 mg OD.'
    eno = '100 IU/kg (1.5 mg/kg), SC, OD.'
    support = 'Hidratim adekuat, ushqyerje dhe pushim; burimi udhëzon shtrim në ICU dhe intubim nëse ka shenja të dështimit respirator, dispne me hipoksemi ose takipne RR >30/min.'
    by_id[repairs.DENGUE]['logicBlocks'] = [b('b9',vitamin+' NËSE TROMBOCITET JANË 10,000 OSE KA GJAKDERDHJE: ≤'),b('b10')]
    by_id[repairs.COVID]['logicBlocks'] = [b('b9',remdesivir.split(' 100 mL 0.9% NaCl, OD,')[0]),b('b10',title='mL 0.9% NaCl, OD, për 45 minuta deri 1 orë. Gjithsej 4-9 doza nga 100 mg OD.'),b('b13',eno+' '+support),b('b14')]
    by_id[repairs.PLEURAL]['logicBlocks'] = [{**b('b2'),'note':'KAPITULLI 4 - PËRSHKRIMI I BARNAVE, AFTËSITË ADMINISTRATIVE'}]
    # Recreate the exact relevant paragraphs under their matching source headings.
    for paragraph in list(rx_source.paragraphs):
        if paragraph.text == 'SOURCE QA 0':
            for value in ['S. '+vitamin,'NËSE TROMBOCITET JANË ≤10,000 OSE KA GJAKDERDHJE:']:
                p = rx_source.add_paragraph(value); paragraph._p.addnext(p._p)
        if paragraph.text == 'SOURCE QA 1':
            for value in ['S. '+remdesivir,'S. '+eno,'S. '+support,'9️⃣ Inj. Enoxaparin sodium 40 mg / 60 mg / 80 mg','🔟 Kujdes mbështetës']:
                p = rx_source.add_paragraph(value); paragraph._p.addnext(p._p)
    rx_path = Path(temp)/'rx.docx';rx_source.save(rx_path)
    result = repairs.build(rx_path,sample,repairs.SOURCE_ID)
    assert result['counts']['changedGuides'] == 3
    fixed = {g['_id']:g for g in result['correctedGuides']}
    dengue_blocks = fixed[repairs.DENGUE]['logicBlocks']
    assert dengue_blocks[0]['items'][0]['sig'] == vitamin
    assert dengue_blocks[1]['condition'] == 'TROMBOCITET JANË ≤10,000 OSE KA GJAKDERDHJE:'
    covid_blocks = fixed[repairs.COVID]['logicBlocks']
    assert covid_blocks[0]['items'][0]['sig'] == remdesivir
    assert not any(x['_key']=='b10' for x in covid_blocks)
    assert next(x for x in covid_blocks if x['_key']=='b13-support')['items'][0]['sig'] == support
    assert repairs.build(rx_path,result['correctedGuides'],repairs.SOURCE_ID)['counts']['changedGuides'] == 0
    try:
        repairs.build(rx_path,sample,'wrong-document')
        raise AssertionError('Wrong source must fail')
    except AssertionError as error:
        assert 'canonical prescription' in str(error)
    for p in rx_source.paragraphs:
        if p.text == 'NËSE TROMBOCITET JANË ≤10,000 OSE KA GJAKDERDHJE:':
            p.text = p.text.replace('≤','≥')
    rx_source.save(rx_path)
    try:
        repairs.build(rx_path,sample,repairs.SOURCE_ID)
        raise AssertionError('Changed source must fail')
    except AssertionError as error:
        assert 'Source paragraph differs' in str(error)
print('Prescription repairs: 74 source headings, exact conditions, complete infusions, separate support, unchanged review status, rollback and idempotence passed.')
