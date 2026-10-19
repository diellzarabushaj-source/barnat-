"""Chapters 15–20 preserve exact source order and stage missing references unapproved."""
import copy
import importlib.util
import tempfile
from pathlib import Path
from docx import Document

s=importlib.util.spec_from_file_location('rx1520',Path(__file__).parents[1]/'scripts/repair-prescription-chapters-15-20.py')
m=importlib.util.module_from_spec(s);s.loader.exec_module(m)
compound=m.engine.connector(['PLUS/MUND TË KOMBINOHET'],{'relation':'conditional'},m.PROFILE['connector_overrides'])
assert compound=={'sourceConnectorLabel':'PLUS/MUND TË KOMBINOHET','condition':'','relation':'plus','selection':'all'}

def block(key,title,sig='Truncated'):
    return {'_key':key,'_type':'prescriptionLogicBlock','relation':'and','selection':'choose-one','sourceConnectorLabel':'DHE','order':1,'items':[{'_key':key+'-item','_type':'sourcePrescriptionItem','kind':'source-note','title':title,'sig':sig,'order':1}]}

with tempfile.TemporaryDirectory() as temp:
    doc=Document();guides=[]
    for chapter,count in m.EXPECTED.items():
        doc.add_heading(f'KAPITULLI {chapter} - QA',2)
        for number in range(1,count+1):
            new=chapter==16 and number in m.NEW_GUIDES
            title=m.NEW_GUIDES[number][1] if new else f'QA {chapter}.{number}'
            doc.add_heading(title,4);doc.add_paragraph('Rx')
            doc.add_paragraph('1️⃣ Tab. QA 10 mg');doc.add_paragraph('S. Full source instruction; ≥30 mL/min; 24 orë.')
            doc.add_paragraph('S. Separate support paragraph.')
            blocks=[block('b1','Tab. QA 10 mg')];blocks[0].update(relation='start',selection='all',sourceConnectorLabel='')
            for key,fragment in m.FRAGMENTS.get((chapter,number),{}).items():blocks.append(block(key,fragment,''))
            for (c,n,key),(old,fixed) in m.NOTE_FIXES.items():
                if (c,n)!=(chapter,number):continue
                doc.add_paragraph('DHE');doc.add_paragraph(fixed);doc.add_paragraph('S. Full source-specific instruction.')
                blocks.append(block(key,old))
            if new:continue
            guides.append({'_id':f'qa-{chapter}-{number}','_rev':'read-revision','title':title,'sourceHeading':title,'chapterNumber':chapter,'orderInChapter':number,'reviewStatus':'review' if number%2 else 'verified','logicBlocks':blocks})
    path=Path(temp)/'source.docx';doc.save(path)
    result=m.build(path,guides)
    assert result['counts']['guidesCompared']==137 and result['counts']['sourceGuides']==140
    assert result['counts']['removedSplitArtifacts']==10 and result['counts']['newReferences']==3
    assert len(result['creates'])==3 and len(result['correctedGuides'])==140
    by_id={g['_id']:g for g in result['correctedGuides']}
    for old in guides:
        new=by_id[old['_id']]
        for field in ('_id','_rev','reviewStatus'):assert new[field]==old[field]
        kept={b['_key'] for b in new['logicBlocks']}
        assert kept=={b['_key'] for b in old['logicBlocks']}-set(m.FRAGMENTS.get((old['chapterNumber'],old['orderInChapter']),{}))
        assert new['logicBlocks'][0]['note']=='S. Separate support paragraph.'
        assert new['logicBlocks'][0]['items'][0]['sig']=='Full source instruction; ≥30 mL/min; 24 orë.'
        for b in new['logicBlocks']:
            assert b['items'][0]['_key']==next(x for x in old['logicBlocks'] if x['_key']==b['_key'])['items'][0]['_key']
    for reference in result['creates']:
        assert reference['reviewStatus']=='source-imported' and '_rev' not in reference
        assert all(i['kind']=='source-note' for b in reference['logicBlocks'] for i in b['items'])
        assert all(not any(k in i for k in ('dose','genericName','strength','route')) for b in reference['logicBlocks'] for i in b['items'])
    repeat=m.build(path,result['correctedGuides'])
    assert repeat['creates']==[] and repeat['changes']==[]
    changed=copy.deepcopy(guides);changed[0]['sourceHeading']='Different source'
    try:m.build(path,changed)
    except AssertionError:pass
    else:raise AssertionError('Source heading changes accepted')
    try:m.build(path,guides,'wrong-source')
    except AssertionError:pass
    else:raise AssertionError('Wrong source accepted')
    doc.add_table(rows=1,cols=1).cell(0,0).text='Unsupported dose table';doc.save(path)
    try:m.build(path,guides)
    except AssertionError:pass
    else:raise AssertionError('Unsupported table omitted')
print('Chapters 15–20: 137 existing guides, 140 source sections, 10 fragments, unapproved references, IDs/status and repeat conversion passed.')
