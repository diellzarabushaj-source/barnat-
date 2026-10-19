"""Preserve all chapters 10–15, including body paragraphs styled as headings."""
import copy
import importlib.util
import tempfile
from pathlib import Path
from docx import Document

s=importlib.util.spec_from_file_location('rx1015',Path(__file__).parents[1]/'scripts/repair-prescription-chapters-10-15.py')
m=importlib.util.module_from_spec(s);s.loader.exec_module(m)

def block(key,title,sig='Truncated'):
    return {'_key':key,'_type':'prescriptionLogicBlock','order':1,'relation':'and','selection':'choose-one','sourceConnectorLabel':'DHE','items':[{'_key':key+'-item','_type':'sourcePrescriptionItem','kind':'source-note','title':title,'sig':sig,'order':1}]}

with tempfile.TemporaryDirectory() as temp:
    doc=Document();guides=[]
    for chapter,count in m.EXPECTED.items():
        doc.add_heading(f'KAPITULLI {chapter} - QA',2)
        for number in range(1,count+1):
            doc.add_heading(f'QA {chapter}.{number}',4)
            embedded=chapter==12 and number in (13,16)
            doc.add_paragraph('Rx',style='Heading 4' if embedded else None)
            first=m.TITLE_FIXES.get((chapter,number,'b1'),'Tab. QA 10 mg')
            if embedded:first='Transfuzion gjaku'
            blocks=[block('b1',first)];blocks[0].update(relation='start',selection='all',sourceConnectorLabel='')
            doc.add_paragraph(('1️⃣ '+first) if not first.startswith('PËR') else first,style='Heading 4' if embedded else None)
            sig='S. Full source instruction; ≥30 mL/min; 4 javë.'
            if embedded and number==13:sig=next(t for (c,t) in m.EMBEDDED_HEADINGS if t.startswith('S. Në thalassemia'))
            doc.add_paragraph(sig,style='Heading 4' if embedded and number==13 else None)
            doc.add_paragraph('S. Separate support paragraph.')
            for key,title in m.FRAGMENTS.get((chapter,number),{}).items():blocks.append(block(key,title,''))
            for (c,n,key),title in m.TITLE_FIXES.items():
                if (c,n)!=(chapter,number) or key=='b1':continue
                doc.add_paragraph('DHE');doc.add_paragraph(title);doc.add_paragraph('S. Full condition-specific instruction.')
                blocks.append(block(key,'DHE' if title.startswith('DHE') else 'PËR'))
            if embedded and number==16:
                doc.add_paragraph('OSE',style='Heading 4');doc.add_paragraph('1️⃣ Exchange transfusion',style='Heading 4')
                doc.add_paragraph(next(t for (c,t) in m.EMBEDDED_HEADINGS if t.startswith('S. Exchange')),style='Heading 4')
                blocks.append(block('b2','Exchange transfusion'))
            guides.append({'_id':f'qa-{chapter}-{number}','_rev':'read-revision','title':f'QA {chapter}.{number}','sourceHeading':f'QA {chapter}.{number}','chapterNumber':chapter,'orderInChapter':number,'reviewStatus':'review' if number%2 else 'verified','logicBlocks':blocks})
    path=Path(temp)/'source.docx';doc.save(path)
    result=m.build(path,guides)
    assert result['counts']['guidesCompared']==149
    assert result['counts']['removedSplitArtifacts']==10
    assert len(m.source_sections(path))==149
    assert len([g for g in result['correctedGuides'] if g['chapterNumber']==12])==18
    for old,new in zip(guides,result['correctedGuides']):
        for field in ('_id','_rev','reviewStatus'):assert old[field]==new[field]
        kept={b['_key'] for b in new['logicBlocks']}
        assert kept=={b['_key'] for b in old['logicBlocks']}-set(m.FRAGMENTS.get((old['chapterNumber'],old['orderInChapter']),{}))
        assert new['logicBlocks'][0]['note']=='S. Separate support paragraph.'
        for b in new['logicBlocks']:
            assert b['items'][0]['sig']!='Truncated'
            if b.get('sourceConnectorLabel')=='DHE':assert b['selection']=='all'
    assert m.build(path,result['correctedGuides'])['changes']==[]
    changed=copy.deepcopy(guides);changed[0]['sourceHeading']='Other source'
    try:m.build(path,changed)
    except AssertionError:pass
    else:raise AssertionError('Source heading changes accepted')
    try:m.build(path,guides,'wrong-source')
    except AssertionError:pass
    else:raise AssertionError('Wrong source accepted')
    for p in doc.paragraphs:
        if p.text=='1️⃣ Exchange transfusion':p.text='Changed embedded heading';break
    doc.save(path)
    try:m.source_sections(path)
    except AssertionError:pass
    else:raise AssertionError('Unrecognized embedded heading accepted')
print('Chapters 10–15: 149 guides, 10 fragments, embedded heading guards, full notes, IDs/status and repeat conversion passed.')
