"""Final source scope: exact instructions, identities, statuses and idempotence."""
import copy,importlib.util,tempfile
from pathlib import Path
from docx import Document
s=importlib.util.spec_from_file_location('rx2022',Path(__file__).parents[1]/'scripts/repair-prescription-chapters-20-22.py')
m=importlib.util.module_from_spec(s);s.loader.exec_module(m)
def block(key,title,sig='Truncated'):
    return {'_key':key,'_type':'prescriptionLogicBlock','relation':'and','selection':'all','order':1,'items':[{'_key':key+'-item','_type':'sourcePrescriptionItem','kind':'source-note','title':title,'sig':sig}]}
with tempfile.TemporaryDirectory() as temp:
    doc=Document();guides=[]
    for chapter,count in m.EXPECTED.items():
        doc.add_heading(f'KAPITULLI {chapter} - QA',2)
        for number in range(1,count+1):
            title=f'QA {chapter}.{number}';doc.add_heading(title,4);doc.add_paragraph('Rx')
            doc.add_paragraph('1️⃣ Tab. QA 10 mg');doc.add_paragraph('S. Full source instruction; ≥30 mL/min; 24 orë.')
            doc.add_paragraph('S. Separate support paragraph.')
            blocks=[block('b1','Tab. QA 10 mg')];blocks[0]['relation']='start'
            for key,fragment in m.FRAGMENTS.get((chapter,number),{}).items():blocks.append(block(key,fragment,''))
            for (c,n,key),(old,fixed) in m.NOTE_FIXES.items():
                if (c,n)!=(chapter,number):continue
                doc.add_paragraph('DHE');doc.add_paragraph(fixed);doc.add_paragraph('S. Full source-specific instruction.')
                blocks.append(block(key,old))
            guides.append({'_id':f'qa-{chapter}-{number}','_rev':'read-revision','title':title,'sourceHeading':title,'chapterNumber':chapter,'orderInChapter':number,'reviewStatus':'review' if number%2 else 'verified','logicBlocks':blocks})
    path=Path(temp)/'source.docx';doc.save(path);result=m.build(path,guides)
    assert result['counts']['guidesCompared']==148 and result['counts']['removedSplitArtifacts']==14
    for old,new in zip(guides,result['correctedGuides']):
        for field in ('_id','_rev','reviewStatus'):assert new[field]==old[field]
        kept={b['_key'] for b in new['logicBlocks']}
        assert kept=={b['_key'] for b in old['logicBlocks']}-set(m.FRAGMENTS.get((old['chapterNumber'],old['orderInChapter']),{}))
        assert new['logicBlocks'][0]['note']=='S. Separate support paragraph.'
        assert new['logicBlocks'][0]['items'][0]['sig']=='Full source instruction; ≥30 mL/min; 24 orë.'
        for b in new['logicBlocks']:
            assert b['items'][0]['_key']==next(x for x in old['logicBlocks'] if x['_key']==b['_key'])['items'][0]['_key']
    assert m.build(path,result['correctedGuides'])['changes']==[]
    altered=copy.deepcopy(guides);altered[0]['sourceHeading']='Different source'
    for payload,source_id in [(altered,m.SOURCE_ID),(guides,'wrong-source')]:
        try:m.build(path,payload,source_id)
        except AssertionError:pass
        else:raise AssertionError('Changed source accepted')
    doc.add_table(rows=1,cols=1).cell(0,0).text='Unsupported dose table';doc.save(path)
    try:m.build(path,guides)
    except AssertionError:pass
    else:raise AssertionError('Unsupported table omitted')
print('Chapters 20–22: 148 guides, 14 split artifacts, full notes, exact identities/status and repeat conversion passed.')
