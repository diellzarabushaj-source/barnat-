#!/usr/bin/env python3
"""Build source-faithful patches from an exported canonical Google Doc.

Usage: python scripts/import-medical-hub-docx.py source.docx existing-topics.json output.json
Requires python-docx. This script never writes to Sanity.
Pass --initialize-placeholders to replace migration-only scaffolds with source
titles and sections. This mode refuses documents containing authored content.
Use --image-assets registry.json to map source image SHA-256 hashes to uploaded
Sanity asset IDs. Missing assets or unsupported drawings stop the import.
"""
import copy
import argparse
import hashlib
import json
import re
from collections import Counter
from docx import Document
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph
from docx.text.run import Run


def key(prefix, index):
    return f"{prefix}-{index}"


def lesson_heading(text, chapter):
    match = re.match(r'^(\d+)\.\s+(.*)', text)
    if match:
        return int(match[1]), match[2]
    match = re.match(r'^(\d+)\.(\d+)\s+[•—–]\s+(.*)', text)
    if match:
        assert int(match[1]) == chapter, f'Heading chapter mismatch: {text}'
        return int(match[2]), match[3]
    return None


def reference_heading(text):
    return text.startswith('Referencat e shënimeve klinike')


def native_marker(p):
    num_pr = p._p.pPr.numPr if p._p.pPr is not None else None
    if num_pr is None:
        return None
    assert num_pr.numId is not None, 'Unresolved inherited Word list'
    num_id = num_pr.numId.val
    if num_id == 0:
        return None
    level = num_pr.ilvl.val if num_pr.ilvl is not None else 0
    numbering = p.part.numbering_part.element
    nums = numbering.xpath(f'./w:num[@w:numId="{num_id}"]')
    assert len(nums) == 1, f'Missing numbering definition {num_id}'
    assert not nums[0].findall(qn('w:lvlOverride')), f'Word list overrides need an explicit resolver: {num_id}'
    abstract_id = nums[0].find(qn('w:abstractNumId')).get(qn('w:val'))
    levels = numbering.xpath(f'./w:abstractNum[@w:abstractNumId="{abstract_id}"]/w:lvl[@w:ilvl="{level}"]')
    assert len(levels) == 1, f'Missing Word list level: {num_id}/{level}'
    definition = levels[0]
    fmt = definition.find(qn('w:numFmt')).get(qn('w:val'))
    pattern = definition.find(qn('w:lvlText')).get(qn('w:val'))
    if fmt == 'bullet':
        assert pattern in '•●·‣▪◦', f'Unsupported Word bullet: {pattern}'
        prefix, kind = pattern + ' ', 'bullet'
    else:
        assert fmt == 'decimal' and level == 0 and pattern == '%1.', f'Unsupported Word number pattern: {fmt}/{pattern}'
        start = definition.find(qn('w:start'))
        value = int(start.get(qn('w:val'))) if start is not None else 1
        # Counting the preceding XML paragraphs also handles a list continuing
        # from another lesson and independent numIds restarting at their start.
        preceding = p._p.xpath(f'preceding::w:p[w:pPr/w:numPr/w:numId[@w:val="{num_id}"]]')
        value += sum((prior.pPr.numPr.ilvl.val if prior.pPr.numPr.ilvl is not None else 0) == level for prior in preceding)
        prefix, kind = f'{value}. ', 'number'
    marks = []
    for name, mark in [('b', 'strong'), ('i', 'em')]:
        props = definition.find(qn('w:rPr'))
        value = props.find(qn('w:' + name)) if props is not None else None
        if value is not None and value.get(qn('w:val'), '1') not in ['0', 'false', 'off']:
            marks.append(mark)
    return {'prefix': prefix, 'kind': kind, 'level': level + 1, 'marks': marks}


def paragraph_block(p, ident):
    assert not p._p.xpath('.//w:drawing | .//w:pict | .//m:oMath'), f'Unsupported image/math at {ident}; import it explicitly instead of discarding it'
    spans, definitions = [], []
    # Iterate XML, including runs inside hyperlinks; Paragraph.runs omits those.
    for elem in p._p:
        runs = [elem] if elem.tag == qn('w:r') else list(elem.iter(qn('w:r'))) if elem.tag == qn('w:hyperlink') else []
        for run_element in runs:
            r = Run(run_element, p)
            if not r.text:
                continue
            marks = []
            for attribute, mark in [('bold', 'strong'), ('italic', 'em'), ('underline', 'underline')]:
                value = getattr(r, attribute)
                if value is None:
                    value = getattr(p.style.font, attribute)
                if value:
                    marks.append(mark)
            if r.font.strike:
                marks.append('strike-through')
            if elem.tag == qn('w:hyperlink'):
                rid = elem.get(qn('r:id'))
                if rid and rid in p.part.rels:
                    link_key = key(ident + '-link', len(definitions))
                    definitions.append({'_type': 'link', '_key': link_key, 'href': p.part.rels[rid].target_ref})
                    marks.append(link_key)
            spans.append({'_type': 'span', '_key': key(ident + '-s', len(spans)), 'text': r.text, 'marks': marks})
    assert ''.join(s['text'] for s in spans) == p.text, f"Unsupported paragraph content at {ident}: {p.text}"
    result = {'_type': 'block', '_key': ident, 'style': 'h3' if p.style.name.startswith('Heading') else 'normal', 'markDefs': definitions, 'children': spans}
    # Keep literal source markers in the stored spans; the renderer consumes them
    # once, preserving the original numbering and bold ranges across run boundaries.
    number = re.match(r'^\s*(\d+)\.\s+', p.text)
    bullet = re.match(r'^(\s*)[•●·‣▪◦]\s+', p.text)
    if number:
        result.update(listItem='number', level=1)
    elif bullet:
        result.update(listItem='bullet', level=2 if len(bullet[1]) >= 2 else 1)
    else:
        marker = native_marker(p)
        if marker:
            result['children'].insert(0, {'_type': 'span', '_key': ident + '-native-marker', 'text': marker['prefix'], 'marks': marker['marks']})
            result.update(listItem=marker['kind'], level=marker['level'])
    return result


def image_block(p, ident, assets, alt, source_document_id='1QN0U5sWSj9GdyNV5oZoZIobmjV0TmCwJD937xzzPgVw'):
    assert not p.text.strip(), f'Mixed inline image/text at {ident}; preserve its inline placement explicitly'
    assert not p._p.xpath('.//w:pict | .//m:oMath'), f'Unsupported drawing/math at {ident}'
    drawings, blips = p._p.xpath('.//w:drawing'), p._p.xpath('.//a:blip')
    assert len(drawings) == len(blips) == 1, f'Unsupported drawing group at {ident}'
    for crop in p._p.xpath('.//a:srcRect'):
        assert all(int(value) == 0 for value in crop.attrib.values()), f'Cropped source image at {ident}; import the crop explicitly'
    rid = blips[0].get(qn('r:embed'))
    assert rid and rid in p.part.rels, f'External/unresolved source image at {ident}'
    blob = p.part.rels[rid].target_part.blob
    digest = hashlib.sha256(blob).hexdigest()
    asset = assets.get(digest)
    assert asset and asset.get('assetId', '').startswith('image-'), f'Upload source image {digest} and supply --image-assets before importing {ident}'
    props = p._p.xpath('.//wp:docPr')
    description = props[0].get('descr') if props else None
    result = {'_type': 'medicalFigure', '_key': ident, 'image': {'_type': 'image', 'asset': {'_type': 'reference', '_ref': asset['assetId']}, 'alt': description or alt}, 'sourceUrl': f'https://docs.google.com/document/d/{source_document_id}/edit'}
    return result, digest


def table_block(table, ident):
    grid = [list(row.cells) for row in table.rows]
    assert len(grid) >= 2, f'Table without a body: {ident}'
    width = len(table.columns)
    assert all(len(row) == width for row in grid), f'Unsupported irregular table: {ident}'
    headers = [cell.text for cell in grid[0]]
    result = {'_type': 'medicalTable', '_key': ident, 'columns': headers, 'rows': []}
    # Rich cells are optional alongside the existing exact plain-text values.
    # They retain merged-cell geometry and support future inline formatting.
    seen = set()
    for ri, row in enumerate(grid):
        rich = []
        for ci, cell in enumerate(row):
            assert not cell._tc.xpath('.//w:tbl'), f'Unsupported nested table at {ident}; do not flatten it'
            if cell._tc in seen:
                continue
            seen.add(cell._tc)
            col_span = 1
            while ci + col_span < width and row[ci + col_span]._tc is cell._tc:
                col_span += 1
            row_span = 1
            while ri + row_span < len(grid) and grid[ri + row_span][ci]._tc is cell._tc:
                row_span += 1
            content = [paragraph_block(p, key(f'{ident}-r{ri}-c{ci}', pi)) for pi, p in enumerate(cell.paragraphs) if p.text or p._p.xpath('.//w:drawing | .//w:pict | .//m:oMath')]
            rich.append({'_type': 'medicalTableCell', '_key': f'{ident}-r{ri}-c{ci}', 'columnIndex': ci, 'colSpan': col_span, 'rowSpan': row_span, 'content': content})
        if ri == 0:
            result['headerCells'] = rich
        else:
            result['rows'].append({'_type': 'medicalTableRow', '_key': f'{ident}-r{ri}', 'cells': [c.text for c in row], 'richCells': rich})
    return result


def initialize_placeholders(doc, existing):
    prepared = copy.deepcopy(existing)
    lookup = {(int(t['chapter']['_ref'].split('-')[-1]), int(t['order'])): t for t in prepared}
    for old in prepared:
        sections = old.get('sections', [])
        assert len(sections) == 1 and sections[0].get('_key') == 'editorial-start', f'Authored sections at {old["_id"]}; initialize explicitly instead'
        content = sections[0].get('content', [])
        assert len(content) == 1 and content[0].get('_key') == 'editorial-note' and content[0].get('_type') == 'clinicalCallout', f'Authored content at {old["_id"]}'
        note = ''.join(span.get('text', '') for block in content[0].get('body', []) for span in block.get('children', []))
        assert note == 'Titulli dhe renditja vijnë nga libri kryesor. Përmbajtja klinike duhet të shkruhet dhe verifikohet në Studio para publikimit.', f'Edited migration note at {old["_id"]}; preserve the author edits'
        old['sections'] = []
    chapter, topic = 0, None
    for element in doc.element.body:
        if element.tag != qn('w:p'):
            continue
        p = Paragraph(element, doc)
        if p.style.name == 'Heading 1':
            match = re.search(r'KAPITULLI\s+(\d+)', p.text)
            chapter, topic = (int(match[1]) if match else 0), None
        elif p.style.name == 'Heading 2':
            if reference_heading(p.text) and topic:
                topic['sections'].append({'_type': 'medicalSection', '_key': f'source-{chapter}-{topic["order"]}-references', 'title': p.text, 'sectionType': 'general'})
                continue
            heading = lesson_heading(p.text, chapter)
            topic = lookup.get((chapter, heading[0])) if heading else None
            if topic:
                topic.update(originalTitle=p.text, title=heading[1])
        elif topic and p.style.name == 'Heading 3' and not p.text.endswith(':'):
            prescription = p.text.startswith('RX') or (topic['title'].startswith('RX') and not p.text.startswith(('BURIME', 'REFERENCA')))
            topic['sections'].append({'_type': 'medicalSection', '_key': f'source-{chapter}-{topic["order"]}-section-{len(topic["sections"])}', 'title': p.text, 'sectionType': 'prescription' if prescription else 'general'})
    return prepared


def build(source_path, existing, initialize=False, image_assets=None, source_document_id='1QN0U5sWSj9GdyNV5oZoZIobmjV0TmCwJD937xzzPgVw'):
    assert re.fullmatch(r'[A-Za-z0-9_-]+', source_document_id), 'Invalid source document ID'
    doc = Document(source_path)
    if initialize:
        existing = initialize_placeholders(doc, existing)
    lookup = {(int(t['chapter']['_ref'].split('-')[-1]), int(t['order'])): t for t in existing}
    chapters = {chapter for chapter, _ in lookup}
    output, counts = [], Counter()
    chapter = 0
    topic = section = None
    source_sequence = []
    emitted_sequence = []
    introduction = []

    def append_section(title, is_introduction=False):
        nonlocal section
        intro_key = f'source-{chapter}-intro'
        matches = [s for s in topic['existing']['sections'] if s['title'] == title and (s['_key'] == intro_key) == is_introduction]
        occurrence = sum(s['title'] == title and (s['_key'] == intro_key) == is_introduction for s in topic['sections'])
        old = matches[occurrence] if occurrence < len(matches) else None
        section = copy.deepcopy(old) if old else {'_type': 'medicalSection', '_key': f"source-{chapter}-{topic['order']}-body-{len(topic['sections'])}", 'title': title, 'sectionType': 'general'}
        if is_introduction:
            section['_key'] = intro_key
        section['content'] = []
        section.pop('summary', None)
        topic['sections'].append(section)

    def append_block(block, source_value):
        source_sequence.append(source_value)
        if topic is None:
            introduction.append((block, source_value))
            counts['introductionBlocks'] += 1
            return
        if section is None:
            append_section('__SOURCE_BODY__')
        section['content'].append(block)
        emitted_sequence.append(source_value)

    for index, element in enumerate(doc.element.body):
        if element.tag == qn('w:p'):
            p = Paragraph(element, doc)
            if p.style.name == 'Heading 1':
                assert not introduction, 'Chapter introduction has no lesson destination'
                match = re.search(r'KAPITULLI\s+(\d+)', p.text)
                chapter = int(match[1]) if match else 0
                topic = section = None
                continue
            if chapter not in chapters:
                continue
            if p.style.name == 'Heading 2':
                if reference_heading(p.text) and topic:
                    source_sequence.append(('p', p.text))
                    append_section(p.text)
                    emitted_sequence.append(('p', section['title']))
                    continue
                heading = lesson_heading(p.text, chapter)
                assert heading, f'Unnumbered lesson at {index}'
                order = heading[0]
                old = lookup[(chapter, order)]
                assert old['originalTitle'] == p.text, f'Topic mismatch: {old["_id"]}'
                topic = {'_id': old['_id'], '_rev': old['_rev'], 'chapterNumber': chapter, 'order': order, 'existing': old, 'sections': []}
                if initialize:
                    topic.update(title=old['title'], originalTitle=old['originalTitle'])
                output.append(topic)
                section = None
                if introduction:
                    append_section('__SOURCE_BODY__', is_introduction=True)
                    section['content'] = [block for block, _ in introduction]
                    emitted_sequence.extend(value for _, value in introduction)
                    introduction.clear()
                    section = None
                counts['topics'] += 1
                continue
            if p._p.xpath('.//w:drawing | .//w:pict | .//m:oMath'):
                block, digest = image_block(p, f'source-image-{index}', image_assets or {}, section['title'] if section else 'Chapter introduction', source_document_id)
                append_block(block, ('image', digest))
                counts['figures'] += 1
                continue
            if not p.text.strip():
                continue
            marker = native_marker(p) if not re.match(r'^\s*(?:\d+\.|[•●·‣▪◦])\s+', p.text) else None
            if topic and p.style.name == 'Heading 3' and any(s['title'] == p.text for s in topic['existing']['sections']):
                source_sequence.append(('p', p.text))
                append_section(p.text)
                emitted_sequence.append(('p', section['title']))
            else:
                block = paragraph_block(p, f'source-p{index}')
                source_text = (marker['prefix'] if marker else '') + p.text
                assert ''.join(s['text'] for s in block['children']) == source_text
                append_block(block, ('p', source_text))
                counts['paragraphs'] += 1
                counts['boldSpans'] += sum('strong' in s['marks'] for s in block['children'])
                if block.get('listItem'):
                    counts['listItems'] += 1
                if marker:
                    counts['nativeListItems'] += 1
        elif element.tag == qn('w:tbl') and chapter in chapters:
            table = Table(element, doc)
            block = table_block(table, f'source-table-{index}')
            source_value = [[c.text for c in row.cells] for row in table.rows]
            assert source_value == [block['columns']] + [row['cells'] for row in block['rows']]
            append_block(block, ('table', source_value))
            counts['tables'] += 1
    assert not introduction, 'Chapter introduction has no lesson destination'
    assert source_sequence == emitted_sequence, 'Source text/table order changed'
    assert len(output) == len(existing), 'Missing or duplicate lesson'
    for topic in output:
        assert all(s['content'] for s in topic['sections']), f'Empty section: {topic["_id"]}'
        keys = [s['_key'] for s in topic['sections']]
        assert len(keys) == len(set(keys)), f'Duplicate section identity: {topic["_id"]}'
        topic.pop('existing')
    return {'sourceSha256': hashlib.sha256(open(source_path, 'rb').read()).hexdigest(), 'sourceDocumentId': source_document_id, 'counts': dict(counts), 'topics': output}


if __name__ == '__main__':
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('source')
    parser.add_argument('existing_path')
    parser.add_argument('output_path')
    parser.add_argument('--initialize-placeholders', action='store_true')
    parser.add_argument('--image-assets')
    parser.add_argument('--source-document-id', default='1QN0U5sWSj9GdyNV5oZoZIobmjV0TmCwJD937xzzPgVw', help='Canonical Google Doc ID, including a different volume')
    args = parser.parse_args()
    assets = json.load(open(args.image_assets)) if args.image_assets else {}
    result = build(args.source, json.load(open(args.existing_path)), args.initialize_placeholders, assets, args.source_document_id)
    with open(args.output_path, 'w') as target:
        json.dump(result, target, ensure_ascii=False, indent=2)
    print(json.dumps(result['counts']))
